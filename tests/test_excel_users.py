"""The people list in Excel: template with everyone in it, fill it, import it back (create + update)."""

import datetime
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from server import excel_users as X  # noqa: E402
from server.core import ServerCore  # noqa: E402
from server.db import check_date  # noqa: E402

PORT = 16350


class ExcelUsersTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.core = c = ServerCore(os.path.join(cls.tmp, "srv"))
        c.config.update(tcp_port=PORT, discovery_port=PORT + 1)
        c.start()
        comp = c.call(c.admin_save_department, name="Compositing")
        c.call(c.admin_save_department, name="Roto", parent_id=comp)
        c.call(c.admin_create_user, must_change=False, username="akash", password="Artist2026",
               display_name="Akash Singh", department="Compositing", designation="Compositing Lead",
               birthday="26-09", joined_on="01-04-2021", employee_id="E100")

    @classmethod
    def tearDownClass(cls):
        cls.core.stop()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_dates(self):
        self.assertEqual(check_date("26-09-1990", "d"), "1990-09-26")
        self.assertEqual(check_date("26/9/90", "d"), "1990-09-26")
        self.assertEqual(check_date("2021-04-01", "d"), "2021-04-01")
        self.assertEqual(check_date("26 Sep", "d", year_optional=True), "--09-26")
        self.assertEqual(check_date("--09-26", "d", year_optional=True), "--09-26")
        self.assertEqual(check_date("29-02", "d", year_optional=True), "--02-29")
        self.assertEqual(check_date("", "d"), "")
        for bad in ("31-02-2020", "26-09", "someday", "26-13-2020"):
            with self.assertRaises(ValueError):
                check_date(bad, "d")

    def test_template_fill_and_import(self):
        from openpyxl import load_workbook
        c = self.core
        path = os.path.join(self.tmp, "people.xlsx")
        X.write_template(path, c.call(c.admin_users), c.call(c.admin_departments), c.call(c.admin_roles))
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["How to fill", "People", "Lists"])
        ws = wb["People"]
        headers = [ws.cell(row=3, column=i).value for i in range(1, len(X.COLUMNS) + 1)]
        self.assertEqual(headers[0], "Username *")
        col = {X.COLUMNS[i][0]: i + 1 for i in range(len(X.COLUMNS))}
        self.assertEqual(ws.cell(row=X.FIRST_ROW, column=col["username"]).value, "akash")   # already listed
        self.assertEqual(ws.cell(row=X.FIRST_ROW, column=col["birthday"]).value, "26-09")
        self.assertTrue(ws.data_validations.dataValidation)                               # dropdowns
        lists = wb["Lists"]
        for dv in ws.data_validations.dataValidation:     # bare range (a leading '=' gives empty dropdowns)
            self.assertFalse(dv.formula1.startswith("="), dv.formula1)
        designations = [lists.cell(row=r, column=3).value for r in range(2, lists.max_row + 1)]
        self.assertIn("Compositing Supervisor", designations)       # every designation, incl. ones made later
        self.assertEqual(ws.cell(row=X.FIRST_ROW, column=1).border.left.color.rgb[-6:], X.LINE)   # visible lines
        # fill it in like a person would
        ws.cell(row=X.FIRST_ROW, column=col["designation"]).value = "Compositing Supervisor"   # change
        new = [("priya.n", "Priya Nair", "Compositing", "Roto", "Roto Artist", "akash", "1990-03-12", "E101"),
               ("rahul", "Rahul Verma", "Lighting", "", "Lighting TD", "", "05-11", "E102"),     # new department
               ("bad name", "Bad", "", "", "", "", "", "")]                                     # refused
        for r, (u, n, d, s, des, rep, bday, emp) in enumerate(new, X.FIRST_ROW + 1):
            for key, value in (("username", u), ("display_name", n), ("department", d), ("section", s),
                               ("designation", des), ("reports_to", rep), ("birthday", bday), ("employee_id", emp)):
                ws.cell(row=r, column=col[key]).value = value or None
        ws.cell(row=X.FIRST_ROW + 1, column=col["joined_on"]).value = datetime.date(2024, 6, 3)
        wb.save(path)
        rows = X.read_rows(path)
        for r in rows:
            r.pop("_row")
        preview = c.call(c.admin_import_users, rows, dry_run=True)
        self.assertEqual(sorted(preview["created"]), ["priya.n", "rahul"])
        self.assertEqual(preview["updated"], ["akash"])
        self.assertIn("Lighting", preview["new_departments"])
        self.assertTrue(any("bad name" in e for e in preview["errors"]))
        self.assertIsNone(c.call(c.db.get_user_by_name, "priya.n"))                  # a dry run saves nothing
        result = c.call(c.admin_import_users, rows)
        self.assertEqual(sorted(result["created"]), ["priya.n", "rahul"])
        self.assertEqual({p["username"] for p in result["passwords"]}, {"priya.n", "rahul"})
        priya = c.call(c.db.get_user_by_name, "priya.n")
        self.assertEqual((priya["department"], priya["section"], priya["birthday"], priya["joined_on"],
                          priya["employee_id"], priya["must_change_pw"]),
                         ("Compositing", "Roto", "1990-03-12", "2024-06-03", "E101", 1))
        self.assertEqual(priya["manager_id"], c.call(c.db.get_user_by_name, "akash")["id"])
        self.assertEqual(c.call(c.db.get_user_by_name, "rahul")["birthday"], "--11-05")
        users = {u["username"]: u for u in c.call(c.admin_users)}
        self.assertEqual(users["akash"]["designation"], "Compositing Supervisor")
        self.assertTrue(c.call(c.db.find_department, "Lighting"))
        # importing the same file again changes nothing
        again = c.call(c.admin_import_users, [r for r in rows if r["username"] != "bad name"], dry_run=True)
        self.assertEqual((again["created"], again["updated"], again["unchanged"]), ([], [], 3))
        # the first-password list
        out = os.path.join(self.tmp, "pw.xlsx")
        X.write_passwords(out, result["passwords"], "Studio")
        self.assertEqual(load_workbook(out).active.cell(row=5, column=2).value, result["passwords"][0]["username"])

    def test_csv_still_works(self):
        path = os.path.join(self.tmp, "old.csv")
        with open(path, "w", encoding="utf-8") as f:
            f.write("username,password,display_name,department\ncat,Artist2026,Cat Iyer,Compositing\n")
        rows = X.read_rows(path)
        self.assertEqual((rows[0]["username"], rows[0]["display_name"], rows[0]["password"]),
                         ("cat", "Cat Iyer", "Artist2026"))


if __name__ == "__main__":
    unittest.main()
